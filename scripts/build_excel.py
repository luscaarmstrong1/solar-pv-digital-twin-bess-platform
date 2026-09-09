from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.models import DISCLAIMER
from src.platform_core import DATA, ROOT, load_results

OUT = ROOT / "Solar_PV_Digital_Twin_BESS_Optimization.xlsx"

def add_df(ws, df, start_row=4, start_col=1, max_rows=None):
    if max_rows:
        df = df.head(max_rows)
    for j, col in enumerate(df.columns, start_col):
        c = ws.cell(start_row, j, col)
        c.fill = PatternFill("solid", fgColor="12372A")
        c.font = Font(color="FFFFFF", bold=True)
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, value in enumerate(row, start_col):
            if hasattr(value, "to_pydatetime"):
                value = value.to_pydatetime().replace(tzinfo=None)
            ws.cell(i, j, value)
    return start_row + len(df) + 1

def style(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    for col in range(1, min(ws.max_column, 12) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws["A1"].font = Font(size=14, bold=True, color="12372A")
    ws["A2"].font = Font(size=9, italic=True, color="263238")

def main():
    r = load_results()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = [
        "01_Dashboard", "02_Inputs", "03_PV_System", "04_Weather", "05_Digital_Twin",
        "06_Performance", "07_Losses", "08_Inverters", "09_Strings", "10_Faults",
        "11_OM", "12_BESS", "13_Dispatch", "14_Energy_Balance", "15_Economics",
        "16_LCOS", "17_Scenarios", "18_Sensitivity", "19_Assumptions", "20_Documentation",
    ]
    for s in sheets:
        ws = wb.create_sheet(s)
        ws["A1"] = s
        ws["A2"] = DISCLAIMER
        style(ws)
    dash = wb["01_Dashboard"]
    kpis = [
        ("PV DC Capacity MWp", r["plant"]["dc_capacity_kwp"] / 1000),
        ("PV AC Capacity MW", r["plant"]["ac_capacity_kw"] / 1000),
        ("Annual PV Generation MWh", r["performance"]["annual_generation_kwh"] / 1000),
        ("Performance Ratio", r["performance"]["performance_ratio"]),
        ("Specific Yield kWh/kWp", r["performance"]["specific_yield_kwh_kwp"]),
        ("Availability", r["performance"]["availability"]),
        ("Self-Consumption Before", r["bess"]["self_consumption_before_pct"]),
        ("Self-Consumption After", r["bess"]["self_consumption_after_pct"]),
        ("Peak Reduction kW", r["bess"]["peak_reduction_kw"]),
        ("Annual Savings BRL", r["economics"]["annual_savings_brl"]),
        ("NPV BRL", r["economics"]["npv_brl"]),
        ("LCOS BRL/kWh", r["economics"]["lcos_brl_kwh"]),
    ]
    add_df(dash, pd.DataFrame(kpis, columns=["KPI", "Value"]), 4, 1)
    dash["D4"] = "Editable BESS Power kW"
    dash["E4"] = r["best_bess"]["bess_kw"]
    dash["D5"] = "Editable BESS Energy kWh"
    dash["E5"] = r["best_bess"]["bess_kwh"]
    dash["D7"] = "Simple CAPEX Formula"
    dash["E7"] = "=E4*650+E5*1200"
    dash["D8"] = "LCOS Cross-check"
    dash["E8"] = "='16_LCOS'!B5"
    dash["D9"] = "NPV Cross-check"
    dash["E9"] = "='15_Economics'!B8"
    weather = pd.read_parquet(DATA / "processed" / "weather.parquet")
    twin = pd.read_parquet(DATA / "processed" / "digital_twin.parquet")
    bess = pd.read_parquet(DATA / "processed" / "bess_dispatch.parquet")
    add_df(wb["04_Weather"], weather[["timestamp", "ghi_w_m2", "dni_w_m2", "dhi_w_m2", "temp_air_c", "wind_speed_m_s"]], 4, 1, 500)
    add_df(wb["05_Digital_Twin"], twin[["timestamp", "expected_power_kw", "measured_power_kw", "relative_error_pct", "residual_kw"]], 4, 1, 1000)
    add_df(wb["07_Losses"], pd.read_csv(DATA / "portfolio" / "loss_tree.csv"), 4, 1)
    add_df(wb["09_Strings"], pd.read_csv(DATA / "portfolio" / "strings.csv"), 4, 1)
    add_df(wb["10_Faults"], pd.read_csv(DATA / "portfolio" / "fault_ground_truth.csv"), 4, 1)
    add_df(wb["11_OM"], pd.read_csv(DATA / "portfolio" / "om_events.csv"), 4, 1)
    add_df(wb["12_BESS"], bess[["timestamp", "soc_pct", "charge_power_kw", "discharge_power_kw", "grid_import_kwh", "grid_export_kwh", "bess_losses_kwh"]], 4, 1, 1000)
    add_df(wb["17_Scenarios"], pd.read_csv(DATA / "portfolio" / "scenarios.csv"), 4, 1)
    add_df(wb["18_Sensitivity"], pd.read_csv(DATA / "portfolio" / "sensitivity.csv"), 4, 1)
    for sheet, data in {
        "02_Inputs": pd.DataFrame([{"Parameter": k, "Value": v, "Classification": "premissa configuravel"} for k, v in r["plant"].items() if not isinstance(v, list)]),
        "03_PV_System": pd.DataFrame([r["plant"]]),
        "06_Performance": pd.DataFrame([r["performance"]]),
        "08_Inverters": pd.DataFrame(r["plant"]["inverters"]),
        "13_Dispatch": bess[["timestamp", "charge_power_kw", "discharge_power_kw", "tariff_brl_kwh"]].head(1000),
        "14_Energy_Balance": bess[["timestamp", "balance_error_kwh", "pv_to_load_kwh", "grid_import_kwh", "grid_export_kwh"]].head(1000),
        "15_Economics": pd.DataFrame(list(r["economics"].items()), columns=["Metric", "Value"]),
        "16_LCOS": pd.DataFrame([["LCOS BRL/kWh", r["economics"]["lcos_brl_kwh"]], ["Formula", "PV lifetime BESS costs / PV useful discharged energy"], ["Classification", "resultado calculado"], ["Useful discharged annual kWh", bess["discharge_energy_kwh"].sum() / 2]], columns=["Metric", "Value"]),
        "19_Assumptions": pd.DataFrame([["Energy tariff offpeak", 0.38], ["Energy tariff base", 0.54], ["Energy tariff peak", 1.05], ["Discount rate", 0.10], ["BESS capex BRL/kWh", 1200], ["BESS capex BRL/kW", 650]], columns=["Assumption", "Value"]),
        "20_Documentation": pd.DataFrame([["Disclaimer", DISCLAIMER], ["Technical sources", "docs/fontes_tecnicas.md"], ["Data catalog", "data/data_catalog.yaml"]], columns=["Topic", "Reference"]),
    }.items():
        add_df(wb[sheet], data, 4, 1)
    perf = wb["06_Performance"]
    perf["J4"] = "Excel PR check"
    perf["K4"] = "=B8"
    econ = wb["15_Economics"]
    econ["D4"] = "Cash flow"
    for i, cf in enumerate(r["economics"].get("cashflows", []) if isinstance(r["economics"], dict) else [], 5):
        econ.cell(i, 4, cf)
    econ["B8"] = r["economics"]["npv_brl"]
    # Charts
    chart = LineChart()
    chart.title = "Expected vs Observed"
    data = Reference(wb["05_Digital_Twin"], min_col=2, max_col=3, min_row=4, max_row=200)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(wb["05_Digital_Twin"], min_col=1, min_row=5, max_row=200))
    wb["05_Digital_Twin"].add_chart(chart, "H4")
    bar = BarChart()
    bar.title = "Loss Tree"
    bar.add_data(Reference(wb["07_Losses"], min_col=2, min_row=4, max_row=15), titles_from_data=True)
    bar.set_categories(Reference(wb["07_Losses"], min_col=1, min_row=5, max_row=15))
    wb["07_Losses"].add_chart(bar, "H4")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        ws.sheet_properties.tabColor = "12372A" if ws.title in ["01_Dashboard", "05_Digital_Twin", "15_Economics"] else "2E7D32"
    wb.save(OUT)
    print(OUT)

if __name__ == "__main__":
    main()
